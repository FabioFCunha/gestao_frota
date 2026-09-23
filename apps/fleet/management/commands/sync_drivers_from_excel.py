from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.fleet.models import AdministrativeUnit, Driver


class Command(BaseCommand):
    help = "Sincroniza o cadastro de motoristas a partir de uma planilha Excel sem apagar histórico."

    def add_arguments(self, parser):
        parser.add_argument(
            "source",
            type=Path,
            help="Caminho da planilha de condutores (.xlsx).",
        )

    def handle(self, *args, **options):
        source = options["source"]

        if not source.is_file():
            raise CommandError(f"Planilha não encontrada: {source}")

        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise CommandError(
                "Instale a dependência openpyxl para importar a planilha."
            ) from error

        workbook = load_workbook(
            source,
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        workbook.close()

        if not rows:
            raise CommandError("A planilha está vazia.")

        columns = {
            str(value).strip(): index
            for index, value in enumerate(rows[0])
            if value is not None
        }

        required_columns = {
            "CnhCondutor",
            "Categoria",
            "Registro",
            "Nome",
            "Unidade",
            "Validade",
            "Telefone",
            "Status",
        }

        missing_columns = required_columns - columns.keys()

        if missing_columns:
            raise CommandError(
                "Colunas obrigatórias ausentes: "
                + ", ".join(sorted(missing_columns))
            )

        drivers = []
        registrations = set()

        for row_number, row in enumerate(rows[1:], start=2):
            name = self.value(row, columns, "Nome")
            registration = self.value(row, columns, "Registro")

            if not name or not registration:
                raise CommandError(
                    f"Linha {row_number}: nome e registro são obrigatórios."
                )

            if registration in registrations:
                raise CommandError(
                    f"Linha {row_number}: registro duplicado ({registration})."
                )

            registrations.add(registration)

            drivers.append(
                {
                    "name": name,
                    "registration": registration,
                    "unit_name": self.value(row, columns, "Unidade"),
                    "phone": self.value(row, columns, "Telefone"),
                    "cnh_number": self.cnh_value(
                        row[columns["CnhCondutor"]]
                    ),
                    "cnh_category": self.value(
                        row,
                        columns,
                        "Categoria",
                    ),
                    "cnh_expiration": self.date_value(
                        row[columns["Validade"]]
                    ),
                    "active": (
                        self.value(
                            row,
                            columns,
                            "Status",
                        ).upper()
                        == "ATIVO"
                    ),
                }
            )

        with transaction.atomic():
            # A planilha representa o cadastro atual.
            # Motoristas ausentes ficam inativos, mas NÃO são excluídos.
            Driver.objects.all().update(active=False)

            units = {}

            created_count = 0
            updated_count = 0

            for item in drivers:
                unit_name = item.pop("unit_name")

                unit = None

                if unit_name:
                    unit = units.get(unit_name)

                    if unit is None:
                        unit, _ = AdministrativeUnit.objects.get_or_create(
                            acronym=unit_name,
                            defaults={"name": unit_name},
                        )
                        units[unit_name] = unit

                registration = item["registration"]

                driver, created = Driver.objects.update_or_create(
                    registration=registration,
                    defaults={
                        "name": item["name"],
                        "unit": unit,
                        "phone": item["phone"],
                        "cnh_number": item["cnh_number"],
                        "cnh_category": item["cnh_category"],
                        "cnh_expiration": item["cnh_expiration"],
                        "active": item["active"],
                    },
                )

                if created:
                    created_count += 1
                else:
                    updated_count += 1

        inactive_count = Driver.objects.filter(active=False).count()

        self.stdout.write(
            self.style.SUCCESS(
                "Sincronização concluída com sucesso."
            )
        )
        self.stdout.write(
            f"Condutores na planilha: {len(drivers)}"
        )
        self.stdout.write(
            f"Motoristas criados: {created_count}"
        )
        self.stdout.write(
            f"Motoristas atualizados: {updated_count}"
        )
        self.stdout.write(
            f"Motoristas inativos no cadastro: {inactive_count}"
        )
        self.stdout.write(
            "Histórico de vínculos e acautelamentos preservado."
        )

    @staticmethod
    def value(row, columns, name):
        value = row[columns[name]]

        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def date_value(value):
        if not value:
            return None

        if isinstance(value, datetime):
            return value.date()

        return value

    @staticmethod
    def cnh_value(value):
        number = str(value).strip()

        return number.zfill(11) if number.isdigit() else number
